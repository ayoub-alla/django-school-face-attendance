import os
from io import BytesIO
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from django.conf import settings


def generate_excel(report, buffer):
    wb = Workbook()
    ws = wb.active
    ws.title = "Rapport de présence"

    # Styles
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1E3A5F', end_color='1E3A5F', fill_type='solid')
    present_fill = PatternFill(start_color='DCFCE7', end_color='DCFCE7', fill_type='solid')
    absent_fill = PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid')
    title_font = Font(name='Calibri', size=14, bold=True, color='1E3A5F')
    subtitle_font = Font(name='Calibri', size=10, color='5B6472')
    normal_font = Font(name='Calibri', size=10)
    bold_font = Font(name='Calibri', size=10, bold=True)
    thin_border = Border(
        left=Side(style='thin', color='E3E6EB'),
        right=Side(style='thin', color='E3E6EB'),
        top=Side(style='thin', color='E3E6EB'),
        bottom=Side(style='thin', color='E3E6EB'),
    )
    center_align = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')

    check = report.attendance_check
    classrooms = check.classrooms.all()
    classroom_names = ', '.join([c.name for c in classrooms]) if classrooms else '—'

    present_count = report.present_students.count()
    absent_count = report.absent_students.count()
    unidentified_count = len(report.unidentified_faces or [])

    row = 1

    # === TITRE ===
    ws.merge_cells('A1:D1')
    ws['A1'] = "Rapport de présence"
    ws['A1'].font = title_font
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    row += 1

    ws.merge_cells('A2:D2')
    ws['A2'] = f"Généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')}"
    ws['A2'].font = subtitle_font
    ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
    row += 2

    # === INFORMATIONS ===
    info_data = [
        ('Date', check.created_at.strftime('%d/%m/%Y à %H:%M')),
        ('Utilisateur', check.performed_by.get_full_name() or check.performed_by.username),
        ('Email', check.performed_by.email or 'Non renseigné'),
        ('Classes', classroom_names),
    ]

    for label, value in info_data:
        ws[f'A{row}'] = label
        ws[f'A{row}'].font = bold_font
        ws[f'A{row}'].alignment = left_align
        ws.merge_cells(f'B{row}:D{row}')
        ws[f'B{row}'] = value
        ws[f'B{row}'].font = normal_font
        ws[f'B{row}'].alignment = left_align
        row += 1

    row += 1

    # === STATISTIQUES ===
    ws[f'A{row}'] = 'Présents'
    ws[f'B{row}'] = 'Absents'
    ws[f'C{row}'] = 'Non identifiés'
    for col in ['A', 'B', 'C']:
        ws[f'{col}{row}'].font = header_font
        ws[f'{col}{row}'].fill = header_fill
        ws[f'{col}{row}'].alignment = center_align
        ws[f'{col}{row}'].border = thin_border
    row += 1

    ws[f'A{row}'] = present_count
    ws[f'B{row}'] = absent_count
    ws[f'C{row}'] = unidentified_count
    for col in ['A', 'B', 'C']:
        ws[f'{col}{row}'].font = Font(name='Calibri', size=14, bold=True)
        ws[f'{col}{row}'].alignment = center_align
        ws[f'{col}{row}'].border = thin_border
    ws[f'A{row}'].fill = present_fill
    ws[f'B{row}'].fill = absent_fill
    ws[f'C{row}'].fill = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid')
    row += 2

    # === TABLEAU DES ÉLÈVES ===
    ws[f'A{row}'] = "Liste des élèves"
    ws[f'A{row}'].font = Font(name='Calibri', size=12, bold=True, color='1B2430')
    row += 1

    # En-têtes du tableau
    headers = ['Nom', 'Classe', 'Statut']
    col_widths = [40, 30, 20]
    for i, (header, width) in enumerate(zip(headers, col_widths), 1):
        col_letter = get_column_letter(i)
        ws[f'{col_letter}{row}'] = header
        ws[f'{col_letter}{row}'].font = header_font
        ws[f'{col_letter}{row}'].fill = header_fill
        ws[f'{col_letter}{row}'].alignment = center_align
        ws[f'{col_letter}{row}'].border = thin_border
        ws.column_dimensions[col_letter].width = width
    row += 1

    # Élèves présents
    for student in report.present_students.all():
        ws[f'A{row}'] = student.name
        ws[f'B{row}'] = student.classroom.name
        ws[f'C{row}'] = 'Présent'
        for col_letter in ['A', 'B', 'C']:
            ws[f'{col_letter}{row}'].font = normal_font
            ws[f'{col_letter}{row}'].alignment = left_align if col_letter != 'C' else center_align
            ws[f'{col_letter}{row}'].border = thin_border
            ws[f'{col_letter}{row}'].fill = present_fill
        ws[f'C{row}'].font = Font(name='Calibri', size=10, bold=True, color='15803D')
        row += 1

    # Élèves absents
    for student in report.absent_students.all():
        ws[f'A{row}'] = student.name
        ws[f'B{row}'] = student.classroom.name
        ws[f'C{row}'] = 'Absent'
        for col_letter in ['A', 'B', 'C']:
            ws[f'{col_letter}{row}'].font = normal_font
            ws[f'{col_letter}{row}'].alignment = left_align if col_letter != 'C' else center_align
            ws[f'{col_letter}{row}'].border = thin_border
            ws[f'{col_letter}{row}'].fill = absent_fill
        ws[f'C{row}'].font = Font(name='Calibri', size=10, bold=True, color='B91C1C')
        row += 1

    # Ajuster la largeur des colonnes
    ws.column_dimensions['A'].width = 35
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 15

    wb.save(buffer)
    return buffer